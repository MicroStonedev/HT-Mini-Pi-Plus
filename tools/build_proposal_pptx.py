# -*- coding: utf-8 -*-
"""群核课题开题汇报 PPT（精简版 8 页，企业视角）。

用法：python tools/build_proposal_pptx.py
输出：开题汇报_群核_空间数据与具身智能训练.pptx
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "开题汇报_群核_空间数据与具身智能训练.pptx")

FONT = "微软雅黑"
BG_DARK, BG_LIGHT = "0F1C2E", "F5F7FA"
PRIMARY, PRIMARY_DK = "17557E", "0E3A57"
ACCENT = "E8833A"
TINT, TINT_SOFT = "DCE7EF", "EAF1F7"
TEXT, MUTED, WHITE = "1B2733", "5C6B7A", "FFFFFF"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.62)


def _set_run(run, size, color=TEXT, bold=False, name=FONT):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, name
    f.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=6, line_spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        for (t, size, color, bold) in runs:
            r = p.add_run()
            r.text = t
            _set_run(r, size, color, bold)
    return box


def add_rect(slide, x, y, w, h, fill=None, line=None, radius=None, shadow=False):
    sp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius is not None:
        sp.adjustments[0] = radius
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(1.0)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if shadow:
        el = sp._element.spPr
        effect = el.makeelement(qn("a:effectLst"), {})
        outer = el.makeelement(qn("a:outerShdw"),
                               {"blurRad": "90000", "dist": "38100", "dir": "5400000",
                                "rotWithShape": "0"})
        clr = el.makeelement(qn("a:srgbClr"), {"val": "1B2733"})
        clr.append(el.makeelement(qn("a:alpha"), {"val": "16000"}))
        outer.append(clr)
        effect.append(outer)
        el.append(effect)
    sp.text_frame.word_wrap = True
    return sp


def page_title(slide, text, sub=None):
    add_text(slide, M, Inches(0.42), Inches(11.6), Inches(0.6),
             [[(text, 30, PRIMARY_DK, True)]])
    if sub:
        add_text(slide, M, Inches(1.04), Inches(11.6), Inches(0.34),
                 [[(sub, 13, MUTED, False)]])


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def new_slide(dark=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(
        BG_DARK if dark else BG_LIGHT)
    return s


# ---- S1 封面
s = new_slide(dark=True)
add_text(s, M, Inches(0.9), Inches(10.5), Inches(0.4),
         [[("THU SIGS AIID × 群核科技 · 企业合作课题 · 开题汇报", 14, TINT, False)]])
add_text(s, M, Inches(1.85), Inches(10.8), Inches(2.0),
         [[("从真实空间到训练数据", 46, WHITE, True)],
          [("空间数据与具身智能训练管线", 20, TINT, False)]], space_after=14)
add_rect(s, M, Inches(3.95), Inches(0.9), Inches(0.045), fill=ACCENT)
add_text(s, M, Inches(4.35), Inches(9.4), Inches(1.4),
         [[("选题：群核科技 方向二「空间数据与具身智能训练」", 15, WHITE, False)],
          [("周期：6 周（2026.10.08 – 11.19）", 15, TINT, False)]], space_after=8)
for i, (n, l) in enumerate([("14 → ≥5", "素材场景到交付场景"),
                            ("1 条", "可复现数据管线"),
                            ("3 个", "留出评测场景")]):
    x = Inches(8.05 + i * 1.66)
    add_text(s, x, Inches(5.2), Inches(1.6), Inches(0.5),
             [[(n, 25, ACCENT, True)]])
    add_text(s, x, Inches(5.77), Inches(1.55), Inches(0.5),
             [[(l, 11, TINT, False)]], line_spacing=1.15)
add_text(s, M, Inches(6.6), Inches(9.0), Inches(0.4),
         [[("汇报人：赵胤泽  赵小刚　·　2026 年 10 月", 13, TINT, False)]])
notes(s, "开场（30秒）：各位老师好。我们组选题是群核科技方向二——空间数据与具身智能训练。一句话概括：把真实空间变成机器人可用的训练数据。今天汇报四件事：课题背景与核心问题、实现路径、我们用到群核的哪些产品和技术、以及需要企业支持讨论的地方。周期六周，请各位老师指正。")

# ---- S2 课题介绍：背景与痛点
s = new_slide()
page_title(s, "课题介绍 ①：背景与痛点", "为什么机器人需要「真实空间」的训练数据")
for i, (t, b, hl) in enumerate([
    ("训练数据与真实世界脱节", "双足机器人训练普遍使用程序化生成的抽象地形，"
     "与真实室内空间的几何分布存在系统性偏差", False),
    ("泛化失效", "训练场上走得稳的机器人，进入没见过的教室、走廊、展厅容易失效——"
     "数据分布覆盖不足是核心原因之一", False),
    ("空间重建数据没有流到训练侧", "群核已具备「真实空间 → 3D 重建」的完整能力"
     "（5 亿+三维场景资产），但「重建结果 → 机器人训练数据」这一段还缺少标准化的开源流程", True),
]):
    y = Inches(1.7) + i * Inches(1.62)
    add_rect(s, M, y, Inches(12.1), Inches(1.42),
             fill=TINT_SOFT if hl else WHITE, shadow=True)
    add_rect(s, M + Inches(0.28), y + Inches(0.28), Inches(0.5), Inches(0.5),
             fill=ACCENT if hl else PRIMARY, radius=0.5)
    add_text(s, M + Inches(0.28), y + Inches(0.36), Inches(0.5), Inches(0.34),
             [[(str(i + 1), 17, WHITE, True)]], align=PP_ALIGN.CENTER)
    add_text(s, M + Inches(1.05), y + Inches(0.2), Inches(10.7), Inches(0.4),
             [[(t, 17, PRIMARY_DK, True)]])
    add_text(s, M + Inches(1.05), y + Inches(0.66), Inches(10.7), Inches(0.62),
             [[(b, 13, TEXT, False)]], line_spacing=1.2)
add_rect(s, M, Inches(6.6), Inches(12.1), Inches(0.66), fill=PRIMARY)
add_text(s, M + Inches(0.3), Inches(6.74), Inches(11.5), Inches(0.4),
         [[("核心痛点：真实空间数据 → 具身训练数据，缺少一条可复现、可验证的开源管线",
            14, WHITE, True)]])
notes(s, "约1分钟：背景三个层次。第一，现在双足机器人训练普遍用程序化地形，和真实室内空间的几何分布有系统性偏差。第二，后果是泛化失效——训练场走得好好的，换个没见过的教室就可能摔。第三点也是课题的立足点：群核已经有完整的真实空间到三维重建能力，五亿多个场景资产，但从重建结果到机器人训练数据这一段，还缺少标准化的开源流程。我们把它总结成核心痛点：缺一条可复现、可验证的管线。这正是FirePlace和智元这些案例之外、高校尺度可以补上的一环。")

# ---- S3 课题介绍：核心问题与实现路径
s = new_slide()
page_title(s, "课题介绍 ②：核心问题与实现路径", "三个核心问题 + 三段式路径（前两段自研）")
add_text(s, M, Inches(1.62), Inches(5.9), Inches(0.4),
         [[("核心解决的三个问题", 16, PRIMARY_DK, True)]])
for i, (h, b) in enumerate([
    ("Q-A 怎么转：3DGS 场景 → 物理碰撞体",
     "3D 高斯泼溅不能直接做碰撞体；对比三条转换路径（SpatialLM 结构化 / "
     "Poisson 网格 / 高度场）的性价比后定案"),
    ("Q-B 转多准：误差是伤害还是增强",
     "重建的噪声、缺洞、尺度偏差对下游训练影响几何？通过误差注入实验给出「可容忍误差带」"),
    ("Q-C 值不值：数据价值如何度量",
     "重建场景训练 vs 程序化地形训练，在留出未见场景上对比跌倒率——用泛化表现度量数据价值"),
]):
    y = Inches(2.08) + i * Inches(1.55)
    add_rect(s, M, y, Inches(5.95), Inches(1.38), fill=WHITE, shadow=True)
    add_text(s, M + Inches(0.24), y + Inches(0.16), Inches(5.5), Inches(0.4),
             [[(h, 14, ACCENT, True)]])
    add_text(s, M + Inches(0.24), y + Inches(0.6), Inches(5.5), Inches(0.68),
             [[(b, 12, MUTED, False)]], line_spacing=1.2)
path = [
    ("① 输入与还原", "素材包 14 个 3D 空间（PLY/SPZ）\n+ 自采视频 → SpatialLM 结构化", PRIMARY),
    ("② 转换与生成（自研核心）", "场景转换器 → USD 碰撞体 + 高度场\nSimReady 数据包（场景+DR+初始分布）", ACCENT),
    ("③ 训练验证", "开源双足本体 + 训练框架\n重建 vs 程序化基线对照；留出集泛化", PRIMARY),
]
add_text(s, Inches(6.9), Inches(1.62), Inches(5.8), Inches(0.4),
         [[("实现路径（六周）", 16, PRIMARY_DK, True)]])
for i, (t, d, c) in enumerate(path):
    y = Inches(2.08) + i * Inches(1.55)
    add_rect(s, Inches(6.9), y, Inches(5.8), Inches(1.38), fill=WHITE, shadow=True)
    add_rect(s, Inches(6.9), y, Inches(5.8), Inches(0.44), fill=c, radius=0.15)
    add_text(s, Inches(7.12), y + Inches(0.06), Inches(5.4), Inches(0.32),
             [[(t, 13, WHITE, True)]])
    add_text(s, Inches(7.12), y + Inches(0.56), Inches(5.4), Inches(0.72),
             [[(d, 11.5, TEXT if c == ACCENT else MUTED, c == ACCENT)]],
             line_spacing=1.18)
    if i < 2:
        ar = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.65),
                                y + Inches(1.4), Inches(0.24), Inches(0.16))
        ar.fill.solid()
        ar.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
        ar.line.fill.background()
        ar.shadow.inherit = False
add_text(s, Inches(6.9), Inches(6.75), Inches(5.8), Inches(0.4),
         [[("前两段自研（本课题工作量），第三段全部使用开源组件", 11.5, MUTED, False)]])
notes(s, "约1分钟：核心问题三个，按管线顺序排列。Q-A怎么转——3DGS场景不能直接做物理碰撞体，要在结构化、网格重建、高度场三条路径里选性价比最高的；Q-B转多准——重建误差对训练是伤害还是天然增强，用误差注入实验给出可容忍误差带；Q-C值不值——这是评测的落点，重建场景训练的策略和程序化地形训练的策略，在留出场景上比跌倒率，泛化差距就是数据价值。右边实现路径三段：输入与还原、转换与生成（自研核心）、训练验证；前两段是我们的工作量，第三段全用开源。")

# ---- S4 使用的群核产品与技术
s = new_slide()
page_title(s, "使用的群核产品与技术", "四个抓手，覆盖输入、结构化、平台能力与社区案例")
items = [
    ("Aholo Reality 素材库", "studio.aholo3d.cn/library",
     "14 个 3D 空间的 PLY / SPZ 原始数据——管线的核心输入；"
     "教室 / 教堂 / 村落等多类型室内外场景", PRIMARY),
    ("SpatialLM（开源）", "github.com/manycore-research",
     "空间结构化主线：点云 → 墙体/地面/家具的语义包围盒与布局；"
     "配套 SpatialLM-Dataset 公开数据集用于复现对照", PRIMARY),
    ("Aholo 开放平台", "labs.aholo3d.cn",
     "空间重建 / 生成 / 编辑 / 理解四项能力的 API 与 SDK——"
     "备选转换路径（视频/照片直接重建）的入口", PRIMARY),
    ("社区案例（课题文档指定）", "FirePlace（CVPR）· 智元机器人等",
     "FirePlace 论文致谢使用 SpatialVerse 场景资产；智元/银河通用/穹彻等"
     "用 SpatialVerse 数据训练——验证「空间数据→具身训练」路线成立", ACCENT),
]
for i, (t, src, b, c) in enumerate(items):
    x = M + (i % 2) * Inches(6.2)
    y = Inches(1.7) + (i // 2) * Inches(2.5)
    add_rect(s, x, y, Inches(5.9), Inches(2.3), fill=WHITE, shadow=True)
    add_rect(s, x, y, Inches(5.9), Inches(0.5), fill=c, radius=0.12)
    add_text(s, x + Inches(0.24), y + Inches(0.08), Inches(5.4), Inches(0.36),
             [[(t, 15, WHITE, True)]])
    add_text(s, x + Inches(0.24), y + Inches(0.62), Inches(5.4), Inches(0.3),
             [[(src, 10.5, MUTED, False)]])
    add_text(s, x + Inches(0.24), y + Inches(0.98), Inches(5.42), Inches(1.2),
             [[(b, 12, TEXT, False)]], line_spacing=1.25)
add_text(s, M, Inches(6.75), Inches(12.1), Inches(0.4),
         [[("并行：报名 10 月群核空间智能 workshop / 黑客松——教程与管线直接衔接，"
            "争取企业技术支持通道", 12, PRIMARY_DK, True)]])
notes(s, "约1分钟：明确汇报我们用到群核的哪些产品技术，四个抓手。第一，Aholo Reality素材库——14个场景的PLY和SPZ原始数据，是管线的核心输入，这个已经下载验证过。第二，SpatialLM开源模型——结构化主线，点云转语义包围盒和布局，配套数据集做复现对照。第三，Aholo开放平台——四项空间能力的API和SDK，是备选转换路径的入口。第四，社区案例——FirePlace论文致谢里写明用了SpatialVerse场景资产，智元等机器人公司也用这批数据训练，这证明路线成立。同时我们会报名10月的workshop和黑客松，教程内容和管线直接衔接。")

# ---- S5 需要支持/讨论的地方
s = new_slide()
page_title(s, "需要支持 / 讨论的地方", "三类：权限与算力 · 数据与技术 · 口径确认")
groups = [
    ("权限与算力（优先确认）", ACCENT, [
        ("算力积分申请入口与额度",
         "权益条款载明企业提供平台算力积分——申请通道、额度、适用范围"
         "（SpatialLM 推理 / RL 训练是否都覆盖）"),
        ("平台 API 权限",
         "Aholo 开放平台四项能力的调用配额与鉴权方式；学生账号的科研使用边界"),
    ]),
    ("数据与技术（影响方案细节）", PRIMARY, [
        ("素材数据规格说明",
         "PLY 点云的坐标系/单位约定（实测教室场景为厘米、Z 轴向上）；"
         "SPZ 格式是否有解析文档或 SDK——我们已逆向 PLY 头部，SPZ 暂以点云为主"),
        ("SpatialLM 复现答疑",
         "issue 区已有复现类讨论（ScanNet 指标、零样本检测、坐标框）；"
         "希望确认 issue 之外的求助渠道与响应方式"),
        ("场景选择建议",
         "14 个素材中哪些更适合作为下游训练输入（几何完整度/尺度）；"
         "以及自采重建的目标空间有什么要求"),
    ]),
    ("口径确认（立项相关）", PRIMARY, [
        ("验收口径与版权",
         "评测报告「第三方可验证」的具体要求（是否需企业提供环境复现）；"
         "按「无需买断」立项、成果开源发布，请确认无冲突"),
    ]),
]
y = Inches(1.55)
for gname, color, items in groups:
    add_rect(s, M, y, Inches(12.1), Inches(0.34), fill=color)
    add_text(s, M + Inches(0.18), y + Inches(0.02), Inches(11.7), Inches(0.28),
             [[(gname, 12.5, WHITE, True)]])
    y += Inches(0.44)
    for h, b in items:
        add_rect(s, M, y, Inches(12.1), Inches(0.62), fill=WHITE)
        add_text(s, M + Inches(0.25), y + Inches(0.02), Inches(3.1), Inches(0.58),
                 [[(h, 12.5, PRIMARY_DK, True)]], anchor=MSO_ANCHOR.MIDDLE)
        add_text(s, M + Inches(3.5), y + Inches(0.02), Inches(8.35), Inches(0.58),
                 [[(b, 11, MUTED, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
        y += Inches(0.70)
    y += Inches(0.10)
notes(s, "约1.5分钟（重点页）：需要群核支持讨论的地方分三类。第一类要优先确认：算力积分的申请入口和额度——权益条款里写了企业提供，但通道和范围需要明确，比如SpatialLM推理和RL训练是不是都覆盖；还有平台API的调用配额。第二类影响方案细节：素材数据规格——我们已经实测发现PLY点云是厘米单位Z轴向上，希望有官方口径确认；SPZ格式有没有解析文档；SpatialLM复现遇到问题时issue区之外有没有直接求助渠道；以及14个素材里哪些更适合做训练输入。第三类是立项口径：评测报告第三方可验证的具体要求，和版权按无需买断开源发布有没有冲突。这些希望今天或本周内确认。")

# ---- S6 计划与交付
s = new_slide()
page_title(s, "六周计划与交付物", "周级里程碑，每周有判定标准")
rows = [
    ("W1", "SpatialLM 复现：结构化素材点云", "≥5 个场景完成结构化，与官方口径对照表成型"),
    ("W2", "场景转换器 v1 → Isaac Lab 场景", "≥3 个场景可加载并用于训练循环"),
    ("W3", "SimReady 格式定型 + 平地基线；中期检查", "数据包要素齐备；基线达开源框架口径"),
    ("W4", "核心对照：重建场景 vs 程序化地形训练", "训练收益与转换路径对比有初步数据"),
    ("W5", "留出集泛化 + 误差注入实验", "跌倒率对比与误差带结论成型"),
    ("W6", "评测报告、仓库整理、期末提交", "可复现文档、代码仓库、演示视频齐备"),
]
add_rect(s, M, Inches(1.6), Inches(1.0), Inches(4.7), fill=PRIMARY)
for j, (w_, m, c) in enumerate(rows):
    y = Inches(1.6) + j * Inches(0.78)
    add_rect(s, M + Inches(1.0), y, Inches(11.1), Inches(0.78),
             fill="FFFFFF" if j % 2 == 0 else TINT_SOFT)
    add_text(s, M + Inches(0.06), y + Inches(0.2), Inches(0.9), Inches(0.36),
             [[(w_, 14, WHITE, True)]], align=PP_ALIGN.CENTER)
    add_text(s, M + Inches(1.2), y + Inches(0.21), Inches(5.0), Inches(0.36),
             [[(m, 12.5, TEXT, False)]])
    add_text(s, M + Inches(6.3), y + Inches(0.12), Inches(5.6), Inches(0.56),
             [[(c, 11.5, PRIMARY_DK, False)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.12)
add_text(s, M, Inches(6.55), Inches(12.1), Inches(0.4),
         [[("核心交付：", 14, PRIMARY_DK, True),
           ("场景转换管线（Docker+一键复现）· SimReady 数据集（≥5 场景）· "
            "可被第三方验证的评测报告", 13, TEXT, False)]])
add_text(s, M, Inches(6.95), Inches(12.1), Inches(0.4),
         [[("风险预案：", 14, PRIMARY_DK, True),
           ("素材/算力/真机每项假设都有降级路径（详见提案第 7 节），主线不受单点失败影响",
            13, TEXT, False)]])
notes(s, "约45秒：六周计划快速过——W1结构化复现、W2转换器、W3格式定型加基线、W4核心对照实验、W5泛化与误差实验、W6交付。每周判定标准都已量化。核心交付三项：转换管线、至少五个场景的SimReady数据集、口径可被第三方验证的评测报告。所有风险假设都配了降级路径，单点失败不影响主线。")

# ---- S7 当前进展
s = new_slide()
page_title(s, "当前进展：已提前完成的验证", "立项阶段即完成输入侧与验证侧的打通")
for i, (h, b, tag) in enumerate([
    ("素材数据实测",
     "下载并解析教室场景点云（197 万点）：确认单位为厘米、Z 轴向上、球谐 DC 颜色可还原；"
     "已写入解析器", "已完成"),
    ("点云 → 仿真场景初步打通",
     "去噪、密度均匀化、家具聚类、高度场提取 → MuJoCo 可行走教室场景"
     "（真实房间尺寸 9.4×8.8 m）；并量化确认该点云中课桌几何不可直接提取", "已完成"),
    ("下游验证器就绪",
     "Mini Pi+ 开源策略已在 MuJoCo 复现（0.51 m/s，与官方口径一致），抗扰边界 80–120 N"
     "——管线第③段随时可接入新场景", "已完成"),
    ("发现的一个数据事实（供讨论）",
     "素材点云呈扫描原点主导分布：97% 点集中于扫描者 0.6 m 内，桌面高度带无课桌结构"
     "——印证了「结构化转换路径」的必要性", "进行中"),
]):
    y = Inches(1.7) + i * Inches(1.3)
    add_rect(s, M, y, Inches(12.1), Inches(1.14), fill=WHITE, shadow=True)
    tag_color = ACCENT if tag == "已完成" else MUTED
    add_rect(s, M + Inches(11.0), y + Inches(0.14), Inches(0.86), Inches(0.34),
             fill=tag_color, radius=0.3)
    add_text(s, M + Inches(11.0), y + Inches(0.19), Inches(0.86), Inches(0.26),
             [[(tag, 10.5, WHITE, True)]], align=PP_ALIGN.CENTER)
    add_text(s, M + Inches(0.26), y + Inches(0.13), Inches(10.5), Inches(0.38),
             [[(h, 14.5, PRIMARY_DK, True)]])
    add_text(s, M + Inches(0.26), y + Inches(0.56), Inches(10.5), Inches(0.5),
             [[(b, 11.5, MUTED, False)]], line_spacing=1.18)
notes(s, "约1分钟：虽然方案是今天定稿，但输入侧和验证侧已经提前打通了。第一，素材下载解析完成，确认了坐标约定。第二，点云到仿真场景的初步管线已经跑通——一个真实尺寸的可行走教室场景。第三，下游验证器就绪：官方策略在MuJoCo复现到0.51米每秒，抗扰边界也测出来了，第③段随时能接新场景。第四条特别值得讨论：我们实测发现素材点云是扫描原点主导的分布，97%的点挤在扫描者脚下0.6米内，桌面高度带根本没有课桌结构——这直接印证了方案里选择「结构化转换路径」而不是直接用点云的必要性，也解释了为什么Q-A这个选型问题是核心。")

# ---- S8 结尾
s = new_slide(dark=True)
add_text(s, M, Inches(2.2), Inches(11.5), Inches(1.0),
         [[("从真实空间到训练数据", 40, WHITE, True)]])
add_text(s, M, Inches(3.4), Inches(11.5), Inches(0.6),
         [[("用群核的空间能力，补上具身智能数据链路的缺失一环", 18, TINT, False)]])
for i, (n, l) in enumerate([("3", "个核心问题"), ("4", "个群核技术抓手"),
                            ("≥5", "个交付场景"), ("6", "周完成交付")]):
    x = M + Inches(i * 3.0)
    add_text(s, x, Inches(4.7), Inches(2.8), Inches(0.7), [[(n, 34, ACCENT, True)]])
    add_text(s, x, Inches(5.45), Inches(2.8), Inches(0.4), [[(l, 13, TINT, False)]])
add_text(s, M, Inches(6.5), Inches(11.5), Inches(0.5),
         [[("谢谢，欢迎讨论与指正", 20, WHITE, True)]])
notes(s, "收尾（15秒）：一句话总结——用群核的空间能力，补上具身智能数据链路缺失的一环。三个核心问题、四个群核技术抓手、至少五个交付场景、六周完成。特别欢迎就第五页的支持清单和我们讨论。谢谢。")

prs.save(OUT)
print("saved:", OUT, os.path.getsize(OUT) // 1024, "KB")
