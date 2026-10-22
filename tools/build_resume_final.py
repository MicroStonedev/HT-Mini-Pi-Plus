# -*- coding: utf-8 -*-
"""实习简历排版美化版：内容与 实习简历2610.docx 完全一致，仅优化布局。

用法：python tools/build_resume_final.py
输出：D:\\研一上课程\\实习求职\\实习简历2610_美化版.docx
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"D:\研一上课程\实习求职\实习简历2610_美化版.docx"
PHOTO = r"D:\研一上课程\实习求职\photo.jpg"

DARK = RGBColor(0x17, 0x3A, 0x5E)    # 深蓝：标题与结构元素
ACCENT = RGBColor(0xB0, 0x5A, 0x1E)  # 琥珀：求职方向等强调
GRAY = RGBColor(0x64, 0x70, 0x7B)    # 次要信息
FONT = "微软雅黑"


def set_font(run, size=9.5, bold=False, color=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rfonts = rPr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rfonts)
    for k in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rfonts.set(qn(k), name)


def no_borders(tbl):
    """按 OOXML schema 顺序移除表格边框（tblBorders 须在 tblLayout/tblLook 之前，
    否则 WPS 会忽略并显示默认虚线框）。"""
    tblPr = tbl._tbl.tblPr
    borders = tblPr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(borders.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "none", qn("w:sz"): "0",
            qn("w:space"): "0", qn("w:color"): "auto"}))
    anchor = None
    for tag in ("w:tblLayout", "w:tblCellMar", "w:tblLook"):
        anchor = tblPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        tblPr.append(borders)
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            if tcPr.find(qn("w:tcBorders")) is None:
                tcb = tcPr.makeelement(qn("w:tcBorders"), {})
                for edge in ("top", "left", "bottom", "right"):
                    tcb.append(tcb.makeelement(qn(f"w:{edge}"), {
                        qn("w:val"): "none", qn("w:sz"): "0",
                        qn("w:space"): "0", qn("w:color"): "auto"}))
                tcPr.append(tcb)


def para(doc, runs, space_after=2, space_before=0, align=None, indent=None,
         line=1.08):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    for text, size, bold, color in runs:
        set_font(p.add_run(text), size, bold, color)
    return p


def section_title(doc, text):
    p = para(doc, [(text, 12, True, DARK)], space_after=4, space_before=9, line=1.0)
    pPr = p._p.get_or_add_pPr()
    borders = pPr.makeelement(qn("w:pBdr"), {})
    borders.append(borders.makeelement(qn("w:bottom"), {
        qn("w:val"): "single", qn("w:sz"): "6",
        qn("w:space"): "3", qn("w:color"): "9DB2C6"}))
    pPr.append(borders)
    return p


def entry_row(doc, time_text, title_text, role_text):
    """项目/教育条目行：时间（固定宽，防止换行）· 标题（加粗）· 角色（右对齐灰）。"""
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = (Cm(3.5), Cm(10.2), Cm(3.9))
    for col, w in zip(t.columns, widths):
        col.width = w
    for i, w in enumerate(widths):
        t.cell(0, i).width = w
    cells = t.rows[0].cells
    for i, (txt, size, bold, color) in enumerate((
            (time_text, 9.5, False, GRAY),
            (title_text, 10, True, DARK),
            (role_text, 9.5, False, GRAY))):
        p = cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        if i == 2:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_font(p.add_run(txt), size, bold, color)
    no_borders(t)
    return t


def bullet(doc, head, body, space_after=2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.08
    pf.left_indent = Cm(0.58)
    pf.first_line_indent = Cm(-0.28)
    set_font(p.add_run("• "), 9.5, False, ACCENT)
    if head:
        set_font(p.add_run(head), 9.5, True, DARK)
    set_font(p.add_run(body), 9.5, False, None)
    return p


doc = Document()
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Cm(1.15)
    sec.left_margin = sec.right_margin = Cm(1.45)

# ---------------- 头部（左：姓名/方向/联系方式；右：证件照）
ht = doc.add_table(rows=1, cols=2)
ht.alignment = WD_TABLE_ALIGNMENT.CENTER
ht.autofit = False
ht.columns[0].width = Cm(14.6)
ht.columns[1].width = Cm(3.1)
ht.rows[0].cells[0].width = Cm(14.6)
ht.rows[0].cells[1].width = Cm(3.1)
no_borders(ht)

left = ht.rows[0].cells[0]
left.vertical_alignment = 1
p0 = left.paragraphs[0]
p0.paragraph_format.space_after = Pt(3)
set_font(p0.add_run("赵 小 刚"), 21, True, DARK)
p1 = left.add_paragraph()
p1.paragraph_format.space_after = Pt(3)
set_font(p1.add_run("求职方向：具身智能 / 机器人算法实习"), 11, True, ACCENT)
p2 = left.add_paragraph()
p2.paragraph_format.space_after = Pt(0)
set_font(p2.add_run("2000.09 丨 中共党员 丨 电话：18851871898 丨 邮箱：1565914951@qq.com"),
         9.5, False, GRAY)

right = ht.rows[0].cells[1]
right.vertical_alignment = 1
rp = right.paragraphs[0]
rp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
if os.path.isfile(PHOTO):
    rp.add_run().add_picture(PHOTO, width=Cm(2.35), height=Cm(3.28))

# ---------------- 教育背景
section_title(doc, "教育背景")
bullet(doc, "2026.09–2028.06　清华大学　电子信息　硕士（在读）：",
       "方向为具身智能与机器人", space_after=3)
bullet(doc, "2018.09–2022.06　东南大学　机械工程　本科：",
       "绩点 3.22/4.0（前 40%）丨 国家励志奖学金 ×2 丨 校级创新体验竞赛一等奖")

# ---------------- 项目经历
section_title(doc, "项目经历")

entry_row(doc, "2026.10 – 至今", "人形机器人行走策略训练与真实场景仿真", "核心成员（2 人）")
para(doc, [("企业合作课题（高擎机电 × 群核科技）核心模块：负责人形机器人 Mini Pi+（20 关节）"
            "强化学习行走策略的 Sim2Real 全链路；前三周完成策略复现、自训练、双引擎验证、"
            "DR 参数实验与误差量化，全部脚本化、结果可复现。", 9.5, False, GRAY)],
     indent=0.58, space_after=4)
bullet(doc, "W1 策略复现与评测基座：",
       "逐行复刻官方策略链路（345 维观测→MLP→20 维动作→PD 控制）；搭建 MuJoCo 无头评测环境，"
       "实测官方策略，速度 0.51 m/s 与官方口径一致")
bullet(doc, "W2 PPO 自训练与双引擎验证：",
       "Isaac Lab 训练 PPO 行走策略（官方奖励配置起步、并行 1024 环境），自训策略导出后 "
       "MuJoCo sim2sim 回放连续行走 ≥60 s、速度 0.5 m/s；与官方策略做双引擎稳定性对照")
bullet(doc, "W3 Domain Randomization 参数实验：",
       "针对\u201c过宽策略保守、过窄真机易摔\u201d的权衡，在摩擦与外力两维多档区间对照训练，"
       "以 sim2sim 存活率确定最终训练参数范围")
bullet(doc, "W3 误差来源量化实验：",
       "单变量对照量化三类因素影响——PD 参数偏移 ±20%–50%、IMU 观测噪声、控制延迟注入，"
       "记录退化曲线并汇总为误差对照表，供真机调试排查")
bullet(doc, "真实场景重建与鲁棒性评测：",
       "无头测试脚本施加周期性推力扰动（0–120 N 四档、方向随机）自动统计摔倒率，确定抗扰边界 "
       "80–120 N；解析激光雷达点云转 MuJoCo 可行走场景", space_after=5)

entry_row(doc, "2022.02 – 2022.06", "机器人紧凑组合关节模块设计（毕业设计）", "独立完成")
bullet(doc, "",
       "面向膝关节假肢测试设计模拟人体下肢的测试平台：分析步态特征，完成躯干–大腿–髋关节结构与"
       "驱动方案设计，关键件通过强度/寿命校核，Adams 动力学仿真验证可行性", space_after=4)

# ---------------- 技能专长
section_title(doc, "技能专长")
bullet(doc, "编程与数据：", "Python（PyTorch、NumPy）、C++、Linux；", space_after=3)
bullet(doc, "机器人仿真：",
       "MuJoCo（策略验证与批量评测，自建无头测试框架）；Isaac Lab（PPO 强化学习训练）", space_after=3)
bullet(doc, "具身智能链路：",
       "熟悉 Sim2Real 方法（Domain Randomization、PD 控制、误差归因）；"
       "了解 ROS 2 推理部署与 RKNN/TensorRT 模型转换链路")

# ---------------- 其他
section_title(doc, "其他")
bullet(doc, "校园经历：", "本科时暑期社会实践项目负责人，调研成果发表于中青在线", space_after=3)
bullet(doc, "自我评价：", "学习能力良好，做事认真负责，具备良好的抗压能力和团队合作能力。",
       space_after=1)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("saved:", OUT, os.path.getsize(OUT) // 1024, "KB")
