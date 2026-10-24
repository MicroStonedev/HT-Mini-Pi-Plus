# -*- coding: utf-8 -*-
"""生成《W3 详细执行步骤》Word 文档"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn

OUT = r"D:\zcodeproject\HT-Mini-Pi-Plus-Robot\docs\W3详细执行步骤.docx"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def para(doc, text, size=10.5, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x8B, 0x2F, 0x2F)
    return p


def simple_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        set_font(r, 10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shd = c._tc.get_or_add_tcPr()
        el = shd.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "1F4E79"})
        shd.append(el)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.rows[i + 1].cells[j]
            c.text = ""
            r = c.paragraphs[0].add_run(str(v))
            set_font(r, 9.5)
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t


doc = Document()
for s in doc.sections:
    s.top_margin, s.bottom_margin = Cm(2), Cm(2)
    s.left_margin, s.right_margin = Cm(2.2), Cm(2.2)

para(doc, "W3 详细执行步骤（10/30–11/04）：交付冲刺与期末演示", 16, True, ACCENT)
para(doc, "")
para(doc, "阶段使命：把 W2 训出的成果变成「能演示、能答辩、能写进简历」的最终交付物。")
para(doc, "时间坐标：期末演示 11/19——W3 收官后仍有约两周缓冲，本周期以质量优先。")
para(doc, "起点资产（W2 交付）：自训 checkpoint（目标 ≥12000 iter）｜消融数据表（扰动/摩擦）｜"
          "sim2sim 对照数字｜全部证据已入 GitHub 仓库。", 10.5)

# ---------------- D10 ----------------
para(doc, "D10 · 10/30（盘点决策日，3–4h）：W2 成果评估与补救决策", 13, True, ACCENT)
para(doc, "步骤 1：最终 checkpoint 大考（30 分钟）", 11, True)
code(doc, "# 云端：用 run1 最终 checkpoint 跑 300 秒回放（同 W1 判定流程）")
code(doc, "ACCEPT_EULA=Y python legged_lab/scripts/play_log.py --task=pi_plus_amp_sym_flat \\")
code(doc, "  --load_run=<run1目录> --checkpoint=model_<最终>.pt --num_envs=1 --headless")
simple_table(doc, ["回放结果", "决策"], [
    ("300s 零跌倒 + 前进速度 ≥0.3 m/s", "达标 → 直接进素材制作（D12）"),
    ("能走但速度/稳定性不足", "D11 补训 3000–5000 iter（调学习率或续训）"),
    ("明显劣化/秒摔", "回退到中间最优 checkpoint，用「次优但可用」交付"),
], [7.5, 8.0])
para(doc, "步骤 2：数据汇总表定稿（把 W2 所有实测填进一张总表）", 11, True)
simple_table(doc, ["指标", "官方策略", "自训策略", "差距分析"], [
    ("平地速度 m/s", "0.51", "___", "一句话原因"),
    ("Isaac 回放 300s", "0 跌倒", "___", ""),
    ("抗扰边界 N", "80–120", "___", ""),
    ("教室场景通过", "6.1m/12s", "___", ""),
], [3.6, 3.4, 3.4, 5.1])
para(doc, "这张表 = 演示 PPT 的核心页 + 简历数字来源 + 面试答案，一天填好全周期受益。", 10.5)

# ---------------- D11 ----------------
para(doc, "D11 · 10/31（弹性日，0–6h）：补训或消融收尾", 13, True, ACCENT)
para(doc, "两种剧本二选一（由 D10 决策决定）：", 11, True)
para(doc, "剧本 A（达标）：不做补训。做消融补充实验（若 W2 未完成 D 组域随机化对比），"
          "或直接休息把时间投给面试线。", 10.5)
para(doc, "剧本 B（补训）：续训 3000–5000 iter（约 4–6 小时）。挂 tmux + 看门狗同 run1 模式。"
          "注意：续训加载用 --resume，改参需记录进误差追踪表。", 10.5)

# ---------------- D12 ----------------
para(doc, "D12 · 11/01（素材制作日，4–5h）：演示视频三件套", 13, True, ACCENT)
para(doc, "素材 1：Isaac 侧自训策略回放视频（云端录制）", 11, True)
para(doc, "开流式窗口（WebRTC）或用录屏方式，录制自训策略 Isaac 行走 ≥60 秒，"
          "含速度指令切换片段（0.3→0.6 m/s 响应）。", 10.5)
para(doc, "素材 2：本地 MuJoCo 对比视频（sim2sim 证据）", 11, True)
code(doc, "# 本地：自训策略在教室场景行走 + 推力扰动片段")
code(doc, "python tools/record_demo.py --policy results/checkpoints/run1_policy.pt \\")
code(doc, "  --model scenes/scene_classroom.xml --push 80")
para(doc, "素材 3：官方 vs 自训对比剪辑（最有说服力的一条）", 11, True)
para(doc, "同视角同场景各录 15 秒，左右拼接——直观展示复现水平。归档全部到 results/videos/。", 10.5)

# ---------------- D13 ----------------
para(doc, "D13 · 11/02（材料更新日，4h）：演示 PPT 与演讲稿 v2", 13, True, ACCENT)
para(doc, "基于现有 11 页期末演示 PPT 升级，重点改四处：", 11, True)
simple_table(doc, ["页", "改动"], [
    ("第 4 页 数据页", "官方/自训对照总表（D10 那张）替换旧数据"),
    ("新增 1 页", "训练曲线：回合长度 17→___ 的爬升 + TensorBoard 截图"),
    ("新增 1 页", "消融实验图表（扰动/摩擦边界图）"),
    ("结尾页", "GitHub 仓库链接 + 三层备份体系一句话（工程能力展示）"),
], [3.5, 12.0])
para(doc, "演讲稿同步更新：把「训练环境搭建中」的诚实口径升级为「自训策略实测 ___ m/s」——"
          "这是从复现到自研的口径跃迁。附二十问按新数据重答（重点：Q 差距原因类问题）。", 10.5)

# ---------------- D14 ----------------
para(doc, "D14 · 11/03（彩排日，3h）：10 分钟全流程计时彩排 ×2", 13, True, ACCENT)
para(doc, "彩排清单：", 11, True)
for line in [
    "① 计时：10 分钟讲完（含视频播放 1.5 分钟）——超时是最常见死法",
    "② 现场演示预案：本地交互版 sim2sim 备好（展示当天的备用方案：纯视频 or 现场交互，二选一）",
    "③ 十问快问快答：队友随机抽问，答不上的记下来当天补",
    "④ 设备检查：HDMI 转接、视频编解码（MP4 用 H.264 通用编码）、PPT 字体嵌入",
]:
    para(doc, line, 10.5)
para(doc, "两次彩排：第一次顺内容，第二次掐时间。录像自看一遍。", 10.5, True)

# ---------------- D15 ----------------
para(doc, "D15 · 11/04（固化收官日，3h）：交付物清单与云资源退场", 13, True, ACCENT)
para(doc, "交付物终检清单：", 11, True)
for line in [
    "☐ GitHub 仓库打 tag v1.0-final，README 更新最终数据",
    "☐ 演示 PPT v2 + 演讲稿 v2 + 对比视频 三件套入仓库",
    "☐ W3 周报（含三周总复盘一段：从复现到自训的完整叙事）",
    "☐ 简历最终版：第一条 bullet 替换自训实测数字",
    "☐ 误差追踪表定稿（15+ 坑，面试工程能力素材）",
    "☐ 云资源退场：实例释放、镜像删除（项目结束，停止一切计费）",
    "☐ 本地备份：仓库 clone 一份到 U 盘/网盘（双保险）",
]:
    para(doc, line, 10.5)

# ---------------- 面试线 ----------------
para(doc, "与面试冲刺并行（每天 2h，W3 重点投入）", 13, True, ACCENT)
simple_table(doc, ["专题", "要求", "自测方法"], [
    ("PPO clip 目标", "手推公式 + 为什么裁剪", "白纸写出损失函数，解释 ε 的作用"),
    ("GAE 优势估计", "λ 偏差-方差权衡", "给队友讲 5 分钟版本"),
    ("AMP 判别器", "与 reward shaping 的区别", "结合本项目：Loss/amp 两条曲线的含义"),
    ("Domain Randomization", "训练分布与部署分布", "结合消融 D 组数据（若做了）"),
    ("项目数字", "全部实测值脱口而出", "模拟面试互问 10 分钟"),
], [4.2, 5.0, 6.3])

para(doc, "速查：W3 最容易翻车的点", 13, True, ACCENT)
simple_table(doc, ["坑", "对策"], [
    ("演示当天视频放不出来", "H.264 编码 + 提前在现场电脑试放 + 备用纯图片版"),
    ("自训效果不达预期心态崩", "诚实口径本身就是加分项：差距分析 > 完美数字"),
    ("彩排超时", "砍视频长度而不是砍讲解——叙事线完整优先"),
    ("云资源忘关持续扣费", "D15 退场清单逐项打勾"),
], [6.0, 9.5])

para(doc, "每日完成判定（打勾制）", 13, True, ACCENT)
for line in [
    "D10：最终 checkpoint 大考结论 + 数据汇总总表填完 ✓",
    "D11：剧本决策执行完（补训完成 or 消融收尾 or 蓄力面试）✓",
    "D12：三条视频素材归档 results/videos/ ✓",
    "D13：PPT v2 + 演讲稿 v2 完成 ✓",
    "D14：两次计时彩排 + 十问无死角 ✓",
    "D15：交付清单全勾 + 云资源清零 ✓",
]:
    para(doc, "☐ " + line, 10.5)

doc.save(OUT)
print("[ok] 已生成", OUT)
